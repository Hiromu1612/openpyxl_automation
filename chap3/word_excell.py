import openpyxl,docx
import os

input_file=r"C:\Users\1612h\Downloads\0src\samples\ch3\meibo.xlsx"
template_file=r"C:\Users\1612h\Downloads\0src\samples\ch3\template-event.docx"

def read_book():
    result=[]
    sheet=openpyxl.load_workbook(input_file).active
    for row in sheet.iter_rows(min_row=2):
        values=[col.value for col in row]
        result.append(values)
        if values[0]==None:
            break
    return result

#顧客の数だけ案内状を作成
for customer in read_book():
    name,post_code,address=customer
    card={
        "住所：▲▲▲":address,
        "〒〇〇〇":post_code,
        "●●様へ":name+"様へ",
    }
    doc=docx.Document(template_file)
    #内容を書き換える
    for p in doc.paragraphs:
        for k,v in card.items():
            p.text=p.text.replace(k,v)
            p.runs[0].font.bold=True #runs[0]は段落の先頭のフォントを取得
    
    save_file="chap3/word_file/"+name+".docx"
    print("saved",save_file)
    if os.path.exists(save_file):
        os.remove(save_file)
    doc.save(save_file)